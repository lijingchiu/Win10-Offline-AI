"""Unit tests for Excel Agent tool functions (openpyxl mode, no Ollama needed)."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent.parent.parent))

import pytest
import openpyxl
from tests.conftest import simple_wb, sales_wb, two_sheet_wb  # noqa: F401


# ── Cleaning ──────────────────────────────────────────────────────────────────

def test_remove_duplicates(simple_wb):
    from agent.tools.cleaning import remove_duplicates
    r = remove_duplicates(simple_wb, sheet="員工資料", columns=["工號"], keep="first")
    assert r["removed_count"] == 1
    assert r["remaining_count"] == 9


def test_strip_whitespace():
    from agent.tools.cleaning import strip_whitespace
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["名稱", "值"])
    ws.append(["  hello  ", "world  "])
    ws.append(["normal", "  text"])
    # All 3 cells with whitespace: "  hello  ", "world  ", "  text"
    r = strip_whitespace(wb)
    assert r["cleaned_count"] == 3
    # Check value is actually stripped
    assert ws.cell(2, 1).value == "hello"


def test_fill_missing_values(simple_wb):
    from agent.tools.cleaning import fill_missing_values
    r = fill_missing_values(simple_wb, sheet="員工資料", columns=["薪資"], fill_value="0", method="value")
    assert r["filled_count"] >= 1


def test_normalize_column_format():
    from agent.tools.cleaning import normalize_column_format
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名"])
    ws.append(["hello world"])
    ws.append(["foo bar"])
    r = normalize_column_format(wb, column="姓名", format_type="upper")
    assert r["error_count"] == 0
    assert ws.cell(2, 1).value == "HELLO WORLD"


def test_split_column():
    from agent.tools.cleaning import split_column
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["地址"])
    ws.append(["台北市,大安區"])
    ws.append(["新北市,板橋區"])
    r = split_column(wb, column="地址", delimiter=",", new_columns=["縣市", "區"])
    assert r["new_columns"] == ["縣市", "區"]
    assert ws.cell(1, 1).value == "縣市"   # header updated
    assert ws.cell(2, 1).value == "台北市"  # data row


def test_merge_columns():
    from agent.tools.cleaning import merge_columns
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓", "名"])
    ws.append(["張", "三"])
    r = merge_columns(wb, columns=["姓", "名"], new_column="全名", separator="")
    assert r["new_column"] == "全名"
    assert ws.cell(2, 1).value == "張三"


# ── Filter/Sort ───────────────────────────────────────────────────────────────

def test_filter_rows_by_condition(simple_wb):
    from agent.tools.filter_sort import filter_rows_by_condition
    r = filter_rows_by_condition(simple_wb, sheet="員工資料", column="部門", operator="=", value="技術")
    assert r["remaining"] == 3  # E002, E005, E007


def test_multi_column_sort(simple_wb):
    from agent.tools.filter_sort import multi_column_sort
    r = multi_column_sort(simple_wb, sheet="員工資料",
                          sort_by=[{"column": "部門", "ascending": True}, {"column": "薪資", "ascending": False}])
    assert "排序" in r["message"]


def test_top_n_rows(simple_wb):
    from agent.tools.filter_sort import top_n_rows
    r = top_n_rows(simple_wb, sheet="員工資料", column="薪資", n=3, largest=True)
    assert r["kept"] == 3


def test_filter_by_range(simple_wb):
    from agent.tools.filter_sort import filter_by_range
    r = filter_by_range(simple_wb, sheet="員工資料", column="薪資", min_value=60000)
    assert r["remaining"] >= 1


def test_filter_by_list(simple_wb):
    from agent.tools.filter_sort import filter_by_list
    r = filter_by_list(simple_wb, sheet="員工資料", column="部門", values=["技術", "人資"])
    assert r["remaining"] >= 1


# ── Stats ─────────────────────────────────────────────────────────────────────

def test_column_sum_avg_count(simple_wb):
    from agent.tools.stats import column_sum_avg_count
    r = column_sum_avg_count(simple_wb, sheet="員工資料", columns=["薪資"])
    assert "總和" in r["message"]


def test_describe_sheet(simple_wb):
    from agent.tools.stats import describe_sheet
    r = describe_sheet(simple_wb, sheet="員工資料")
    assert r["rows"] == 10
    assert "工號" in r["columns"]


def test_group_by_stats(sales_wb):
    from agent.tools.stats import group_by_stats
    r = group_by_stats(sales_wb, sheet="銷售", group_column="地區", agg_column="銷售額", agg_func="sum")
    assert r["groups"] == 2
    assert "群組統計" in sales_wb.sheetnames


def test_conditional_sum(simple_wb):
    from agent.tools.stats import conditional_sum
    r = conditional_sum(simple_wb, sheet="員工資料", sum_column="薪資",
                        condition_column="部門", condition_value="技術")
    assert r["count"] == 3
    assert r["sum"] > 0


def test_value_counts(simple_wb):
    from agent.tools.stats import value_counts
    r = value_counts(simple_wb, sheet="員工資料", column="部門")
    assert "業務" in r["message"]


# ── Pivot ─────────────────────────────────────────────────────────────────────

def test_create_pivot_table(sales_wb):
    from agent.tools.pivot import create_pivot_table
    r = create_pivot_table(sales_wb, sheet="銷售", rows=["地區"], values="銷售額", aggfunc="sum")
    assert "樞紐表" in sales_wb.sheetnames
    assert r["rows_count"] >= 1


def test_cross_tabulation(sales_wb):
    from agent.tools.pivot import cross_tabulation
    r = cross_tabulation(sales_wb, sheet="銷售", row_column="地區", col_column="月份")
    assert "交叉表" in sales_wb.sheetnames


# ── Lookup ────────────────────────────────────────────────────────────────────

def test_find_diff_between_sheets(two_sheet_wb):
    from agent.tools.lookup import find_diff_between_sheets
    r = find_diff_between_sheets(two_sheet_wb, sheet1="Sheet1", sheet2="Sheet2", key_column="ID")
    assert r["only_in_sheet1"] == 2   # A001, A002 only in Sheet1
    assert r["only_in_sheet2"] == 2   # A006, A007 only in Sheet2


def test_vlookup_equivalent(two_sheet_wb):
    from agent.tools.lookup import vlookup_equivalent
    r = vlookup_equivalent(two_sheet_wb, sheet="Sheet1", key_column="ID",
                           lookup_sheet="Sheet2", lookup_key="ID", lookup_value="Name",
                           new_column_name="Sheet2_Name")
    assert r["matched"] >= 1


# ── IO Convert ───────────────────────────────────────────────────────────────

def test_sheet_to_csv(simple_wb):
    from agent.tools.io_convert import sheet_to_csv
    r = sheet_to_csv(simple_wb, sheet="員工資料")
    assert "csv_content" in r
    assert "工號" in r["csv_content"]


def test_csv_to_sheet():
    from agent.tools.io_convert import csv_to_sheet
    wb = openpyxl.Workbook()
    wb.active.append(["dummy"])
    csv_text = "名稱,金額\n蘋果,100\n香蕉,80"
    r = csv_to_sheet(wb, csv_content=csv_text, sheet_name="水果")
    assert "水果" in wb.sheetnames
    assert r["rows"] == 2


def test_merge_multiple_sheets():
    from agent.tools.io_convert import merge_multiple_sheets
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "A"
    ws1.append(["X", "Y"])
    ws1.append([1, 2])
    ws2 = wb.create_sheet("B")
    ws2.append(["X", "Y"])
    ws2.append([3, 4])
    r = merge_multiple_sheets(wb, sheets=["A", "B"], new_sheet="合併")
    assert "合併" in wb.sheetnames
    assert r["total_rows"] == 2


def test_split_sheet_by_column(simple_wb):
    from agent.tools.io_convert import split_sheet_by_column
    r = split_sheet_by_column(simple_wb, sheet="員工資料", column="部門")
    assert len(r["sheets_created"]) == 3  # 業務、技術、人資


# ── Registry ─────────────────────────────────────────────────────────────────

def test_registry_loads_all_tools():
    from agent.registry import all_tool_descriptors
    tools = all_tool_descriptors()
    assert len(tools) >= 38


def test_registry_bm25_retrieval():
    from agent.registry import retrieve_tools
    results = retrieve_tools("刪除重複列", k=5)
    names = [r["name"] for r in results]
    assert "remove_duplicates" in names


def test_registry_english_query():
    from agent.registry import retrieve_tools
    results = retrieve_tools("group by sum statistics", k=5)
    names = [r["name"] for r in results]
    assert any(n in names for n in ["group_by_stats", "group_and_aggregate", "create_pivot_table"])


def test_registry_chart_query():
    from agent.registry import retrieve_tools
    results = retrieve_tools("畫折線圖", k=3)
    names = [r["name"] for r in results]
    assert "insert_line_chart" in names
