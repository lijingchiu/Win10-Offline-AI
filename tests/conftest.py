"""Shared fixtures for unit tests."""
import io
import sys
from pathlib import Path

import pytest
import openpyxl

# Ensure project root on path
sys.path.insert(0, str(Path(__file__).parent.parent))


@pytest.fixture
def simple_wb():
    """Workbook with one sheet: header + 10 rows of employee data."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "員工資料"
    headers = ["工號", "姓名", "部門", "薪資", "在職"]
    ws.append(headers)
    rows = [
        ["E001", "張三", "業務", 55000, "是"],
        ["E002", "李四", "技術", 72000, "是"],
        ["E001", "張三", "業務", 55000, "是"],   # duplicate
        ["E003", "王五", "業務", 48000, "是"],
        ["E004", "趙六", "人資", 60000, "否"],
        ["E005", "陳七", "技術", 85000, "是"],
        ["E006", "林八", "業務", 52000, "是"],
        ["E007", "黃九", "技術", 78000, "是"],
        ["E008", "吳十", "人資", 63000, "是"],
        ["E009", "周十一", "業務", None, "是"],   # missing salary
    ]
    for r in rows:
        ws.append(r)
    return wb


@pytest.fixture
def sales_wb():
    """Workbook with monthly sales data."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "銷售"
    ws.append(["月份", "地區", "銷售額"])
    data = [
        ["2024-01", "北部", 120000],
        ["2024-01", "南部", 95000],
        ["2024-02", "北部", 135000],
        ["2024-02", "南部", 102000],
        ["2024-03", "北部", 98000],
        ["2024-03", "南部", 88000],
    ]
    for r in data:
        ws.append(r)
    return wb


@pytest.fixture
def two_sheet_wb():
    """Workbook with two sheets for comparison tests."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["ID", "Name"])
    for i in range(1, 6):
        ws1.append([f"A{i:03d}", f"Name{i}"])
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["ID", "Name"])
    for i in range(3, 8):
        ws2.append([f"A{i:03d}", f"Name{i}"])
    return wb
