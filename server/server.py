"""
WinLLM Local File Processing Server
Port 8765 — serves frontend, document extraction, and Excel Agent.
"""
import io
import json
import mimetypes
import os
import re
import sys
import uuid
import threading
from datetime import datetime
from http.server import HTTPServer, BaseHTTPRequestHandler
from pathlib import Path
from urllib.parse import urlparse

# ── Path resolution: works both in production (C:\WinLLM\) and dev (repo root) ──
_ROOT = Path(__file__).parent
if not (_ROOT / "agent").exists() and (_ROOT.parent / "agent").exists():
    _ROOT = _ROOT.parent
sys.path.insert(0, str(_ROOT))

FRONTEND_DIR = _ROOT / "frontend"
SESSIONS_DIR = _ROOT / "sessions"
BACKUPS_DIR  = _ROOT / "backups"
PORT = 8765

SESSIONS_DIR.mkdir(parents=True, exist_ok=True)
BACKUPS_DIR.mkdir(parents=True, exist_ok=True)

# ── Global state ──────────────────────────────────────────────────────────────
SESSIONS: dict[str, dict] = {}
DOWNLOAD_TOKENS: dict[str, dict] = {}  # token → {path, filename, content_type}
_lock = threading.Lock()

# ── Agent modules (optional — graceful degradation if not installed) ──────────
try:
    from agent import registry, llm_client, prompts, excel_io, runner
    import agent.llm_client as _lc
    AGENT_ENABLED = True
except ImportError as _e:
    AGENT_ENABLED = False
    print(f"[WinLLM] Agent 模組未載入（文件問答仍可使用）: {_e}")


# ─────────────────────────────────────────────────────────────────────────────
# Document extractors (unchanged from original)
# ─────────────────────────────────────────────────────────────────────────────

def extract_pdf(data: bytes) -> str:
    try:
        import fitz
        doc = fitz.open(stream=data, filetype="pdf")
        return "\n\n".join(page.get_text() for page in doc)
    except ImportError:
        return "[錯誤：未安裝 pymupdf]"
    except Exception as e:
        return f"[PDF 解析錯誤: {e}]"


def extract_docx(data: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        tables_text = []
        for table in doc.tables:
            for row in table.rows:
                tables_text.append(" | ".join(c.text for c in row.cells))
        return "\n".join(paragraphs + tables_text)
    except ImportError:
        return "[錯誤：未安裝 python-docx]"
    except Exception as e:
        return f"[DOCX 解析錯誤: {e}]"


def extract_xlsx(data: bytes) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            parts.append(f"=== 工作表: {sheet_name} ===")
            for row in ws.iter_rows(values_only=True):
                row_vals = [str(v) if v is not None else "" for v in row]
                if any(v.strip() for v in row_vals):
                    parts.append(" | ".join(row_vals))
        return "\n".join(parts)
    except ImportError:
        return "[錯誤：未安裝 openpyxl]"
    except Exception as e:
        return f"[XLSX 解析錯誤: {e}]"


def extract_pptx(data: bytes) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            texts = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            if texts:
                slides.append(f"--- 第 {i} 張投影片 ---\n" + "\n".join(texts))
        return "\n\n".join(slides)
    except ImportError:
        return "[錯誤：未安裝 python-pptx]"
    except Exception as e:
        return f"[PPTX 解析錯誤: {e}]"


# ─────────────────────────────────────────────────────────────────────────────
# Multipart parser (unchanged)
# ─────────────────────────────────────────────────────────────────────────────

def parse_multipart(data: bytes, boundary: bytes) -> dict[str, tuple[str, bytes]]:
    results = {}
    delimiter = b"--" + boundary
    for part in data.split(delimiter)[1:]:
        if part in (b"--\r\n", b"--"):
            continue
        header_end = part.find(b"\r\n\r\n")
        if header_end < 0:
            continue
        headers_raw = part[:header_end].decode("utf-8", errors="replace")
        body = part[header_end + 4:]
        if body.endswith(b"\r\n"):
            body = body[:-2]
        name = filename = None
        for line in headers_raw.splitlines():
            if "Content-Disposition" in line:
                for seg in line.split(";"):
                    seg = seg.strip()
                    if seg.startswith("name="):
                        name = seg.split("=", 1)[1].strip('"')
                    elif seg.startswith("filename="):
                        filename = seg.split("=", 1)[1].strip('"')
        if name:
            results[name] = (filename, body)
    return results


# ─────────────────────────────────────────────────────────────────────────────
# Session helpers
# ─────────────────────────────────────────────────────────────────────────────

def _ensure_session(sid: str, **kwargs) -> dict:
    with _lock:
        if sid not in SESSIONS:
            SESSIONS[sid] = {
                "workbook_path": None,
                "filename": "",
                "history": [],
                "pending_confirm": None,
                "created_at": datetime.now().isoformat(),
                "last_active": datetime.now().isoformat(),
            }
        SESSIONS[sid].update(kwargs)
        SESSIONS[sid]["last_active"] = datetime.now().isoformat()
        return SESSIONS[sid]


def _update_history(sid: str, user_msg: str, assistant_msg: str) -> None:
    with _lock:
        if sid not in SESSIONS:
            return
        h = SESSIONS[sid]["history"]
        h.append({"role": "user",      "content": user_msg})
        h.append({"role": "assistant", "content": assistant_msg})
        SESSIONS[sid]["history"] = h[-12:]  # keep last 6 rounds


def _gc_sessions() -> None:
    """Remove sessions inactive for > 2 hours."""
    import shutil
    now = datetime.now()
    with _lock:
        dead = [
            sid for sid, s in SESSIONS.items()
            if (now - datetime.fromisoformat(s["last_active"])).seconds > 7200
        ]
    for sid in dead:
        try:
            shutil.rmtree(str(SESSIONS_DIR / sid), ignore_errors=True)
        except Exception:
            pass
        with _lock:
            SESSIONS.pop(sid, None)


# Backup index helpers

def _backup_index_path() -> Path:
    return BACKUPS_DIR / "index.json"


def _load_backup_index() -> list:
    p = _backup_index_path()
    if not p.exists():
        return []
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return []


def _save_backup_index(entries: list) -> None:
    _backup_index_path().write_text(
        json.dumps(entries, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def _register_backup(backup_path: str, original_filename: str, tool_name: str, tool_zh: str, sid: str) -> str:
    bid = str(uuid.uuid4())[:8]
    entries = _load_backup_index()
    entries.insert(0, {
        "id": bid,
        "timestamp": datetime.now().isoformat(timespec="seconds"),
        "original_filename": original_filename,
        "tool": tool_name,
        "tool_zh": tool_zh,
        "backup_path": str(Path(backup_path).relative_to(_ROOT)),
        "session_id": sid,
    })
    _save_backup_index(entries)
    return bid


def _format_result(result: dict) -> str:
    return result.get("message") or result.get("answer") or "操作完成"


# ─────────────────────────────────────────────────────────────────────────────
# HTTP Handler
# ─────────────────────────────────────────────────────────────────────────────

class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        pass

    # Allowed origins (localhost only). 'null' (file://) is intentionally NOT allowed —
    # any local HTML file would otherwise be able to read the API.
    _ALLOWED_ORIGINS = {
        "http://localhost:8765", "http://127.0.0.1:8765",
        "http://localhost",      "http://127.0.0.1",
    }

    def _cors(self):
        origin = self.headers.get("Origin", "")
        # Reflect only allowed origins; otherwise omit the header (browser will block).
        if origin in self._ALLOWED_ORIGINS:
            self.send_header("Access-Control-Allow-Origin", origin)
            self.send_header("Vary", "Origin")
        self.send_header("Access-Control-Allow-Methods", "GET,POST,DELETE,OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")

    def do_OPTIONS(self):
        self.send_response(200)
        self._cors()
        self.end_headers()

    # ── GET ──────────────────────────────────────────────────────────────────

    def do_GET(self):
        path = urlparse(self.path).path

        if path == "/api/status":
            self._json({"status": "ok", "version": "2.0.0", "agent": AGENT_ENABLED})
            return
        if path == "/api/agent/tools":
            self._handle_agent_tools()
            return
        if path.startswith("/api/agent/download/"):
            self._handle_agent_download(path.split("/")[-1])
            return
        if path.startswith("/api/agent/sessions/"):
            self._handle_agent_session(path.split("/")[-1])
            return
        if path == "/api/agent/backups":
            self._handle_backups_list()
            return

        # Static files — resolve and verify the path stays inside FRONTEND_DIR
        rel = path.lstrip("/") or "index.html"
        try:
            fpath = (FRONTEND_DIR / rel).resolve()
            frontend_root = FRONTEND_DIR.resolve()
            # Python 3.9+: Path.is_relative_to. Use the manual check for portability.
            inside = str(fpath).startswith(str(frontend_root) + os.sep) or fpath == frontend_root
        except Exception:
            inside = False
        if inside and fpath.exists() and fpath.is_file():
            mime, _ = mimetypes.guess_type(str(fpath))
            data = fpath.read_bytes()
            self.send_response(200)
            self.send_header("Content-Type", mime or "application/octet-stream")
            self.send_header("Content-Length", len(data))
            self._cors()
            self.end_headers()
            self.wfile.write(data)
        else:
            self._not_found()

    # ── POST ─────────────────────────────────────────────────────────────────

    def do_POST(self):
        path = urlparse(self.path).path
        if path == "/api/extract":
            self._handle_extract()
        elif path == "/api/agent/chat":
            self._handle_agent_chat()
        elif path == "/api/agent/confirm":
            self._handle_agent_confirm()
        elif path == "/api/agent/backups/restore":
            self._handle_backup_restore()
        else:
            self._not_found()

    # ── DELETE ───────────────────────────────────────────────────────────────

    def do_DELETE(self):
        path = urlparse(self.path).path
        if path.startswith("/api/agent/sessions/"):
            self._handle_session_delete(path.split("/")[-1])
        elif path.startswith("/api/agent/backups/"):
            self._handle_backup_delete(path.split("/")[-1])
        else:
            self._not_found()

    # ── Document extraction (original, unchanged) ─────────────────────────────

    def _handle_extract(self):
        ct = self.headers.get("Content-Type", "")
        body = self.rfile.read(int(self.headers.get("Content-Length", 0)))
        filename = ""
        file_data = b""
        if "multipart/form-data" in ct:
            boundary = ct.split("boundary=")[-1].encode()
            parts = parse_multipart(body, boundary)
            if "file" in parts:
                filename, file_data = parts["file"]
        else:
            file_data = body
        if not file_data:
            self._json({"error": "no file data"}, 400)
            return
        ext = Path(filename or "").suffix.lower()
        if ext == ".pdf":
            text = extract_pdf(file_data)
        elif ext == ".docx":
            text = extract_docx(file_data)
        elif ext in (".xlsx", ".xls"):
            text = extract_xlsx(file_data)
        elif ext in (".pptx", ".ppt"):
            text = extract_pptx(file_data)
        else:
            text = file_data.decode("utf-8", errors="replace")
        self._json({"filename": filename, "text": text, "chars": len(text)})

    # ── Agent: list tools ─────────────────────────────────────────────────────

    def _handle_agent_tools(self):
        if not AGENT_ENABLED:
            self._json({"error": "Agent 模組未啟用"}, 503)
            return
        tools = registry.all_tool_descriptors()
        self._json({"tools": tools, "count": len(tools)})

    # ── Agent: session info ───────────────────────────────────────────────────

    @staticmethod
    def _is_safe_sid(sid: str) -> bool:
        return isinstance(sid, str) and bool(re.fullmatch(r"[A-Za-z0-9_\-]{1,64}", sid))

    def _handle_agent_session(self, sid: str):
        if not self._is_safe_sid(sid):
            self._json({"error": "invalid session_id"}, 400)
            return
        s = SESSIONS.get(sid)
        if not s:
            self._json({"session_id": sid, "active": False})
            return
        self._json({
            "session_id":    sid,
            "active":        True,
            "filename":      s.get("filename", ""),
            "history_turns": len(s.get("history", [])) // 2,
            "last_active":   s.get("last_active", ""),
        })

    def _handle_session_delete(self, sid: str):
        if not self._is_safe_sid(sid):
            self._json({"error": "invalid session_id"}, 400)
            return
        import shutil
        with _lock:
            SESSIONS.pop(sid, None)
        shutil.rmtree(str(SESSIONS_DIR / sid), ignore_errors=True)
        self._json({"status": "deleted", "session_id": sid})

    # ── Agent: chat (SSE) ─────────────────────────────────────────────────────

    def _handle_agent_chat(self):
        if not AGENT_ENABLED:
            self.send_response(200)
            self.send_header("Content-Type", "text/event-stream; charset=utf-8")
            self._cors()
            self.end_headers()
            self._emit({"type": "error", "text": "Agent 模組未啟用，請確認安裝完整"})
            self._emit({"type": "done"})
            return

        ct = self.headers.get("Content-Type", "")
        body = self.rfile.read(int(self.headers.get("Content-Length", 0)))

        # Set SSE headers immediately so browser shows activity
        self.send_response(200)
        self.send_header("Content-Type", "text/event-stream; charset=utf-8")
        self.send_header("Cache-Control", "no-cache")
        self.send_header("Connection", "keep-alive")
        self._cors()
        self.end_headers()

        try:
            self._agent_chat_logic(ct, body)
        except Exception as e:
            try:
                self._emit({"type": "error", "text": f"伺服器錯誤: {e}"})
            except Exception:
                pass
        finally:
            try:
                self._emit({"type": "done"})
            except Exception:
                pass

    def _agent_chat_logic(self, ct: str, body: bytes):
        _gc_sessions()

        # ── Parse request ──
        if "multipart/form-data" in ct:
            boundary = ct.split("boundary=")[-1].encode()
            parts = parse_multipart(body, boundary)
            def field(k): return parts.get(k, (None, b""))[1].decode("utf-8", errors="replace")
            file_part = parts.get("file")
        else:
            try:
                data = json.loads(body)
            except Exception:
                data = {}
            parts = data
            def field(k): return str(data.get(k, ""))
            file_part = None

        raw_sid = field("session_id") or str(uuid.uuid4())
        # Reject anything that isn't a safe identifier (prevent path traversal / injection
        # via session_id used as a directory name).
        sid = raw_sid if re.fullmatch(r"[A-Za-z0-9_\-]{1,64}", raw_sid) else str(uuid.uuid4())
        message = field("message").strip()

        if not message:
            self._emit({"type": "error", "text": "訊息不能為空"})
            return

        # ── File upload ──
        if file_part:
            filename, file_data = file_part
            if file_data:
                session_dir = SESSIONS_DIR / sid
                session_dir.mkdir(parents=True, exist_ok=True)
                working_path = session_dir / "working.xlsx"
                working_path.write_bytes(file_data)
                _ensure_session(sid, filename=filename or "工作簿.xlsx",
                                workbook_path=str(working_path))
                self._emit({"type": "file_loaded", "filename": filename})

        session = _ensure_session(sid)
        if not session.get("workbook_path"):
            self._emit({"type": "error", "text": "請先上傳 Excel 檔案（.xlsx）"})
            return

        self._emit({"type": "thinking", "text": "正在分析您的需求…"})

        # ── Load workbook ──
        try:
            wb, mode = excel_io.open_workbook(session["workbook_path"])
        except Exception as e:
            self._emit({"type": "error", "text": f"無法開啟工作簿: {e}"})
            return

        wb_info = excel_io.get_wb_info(wb, session.get("filename", "工作簿"))
        self._emit({"type": "mode", "mode": mode})

        # ── BM25 retrieve top-5 tools ──
        top_tools = registry.retrieve_tools(message, k=5)

        # ── Build prompt + call LLM ──
        messages = prompts.build_prompt(top_tools, wb_info, session.get("history", []), message)
        try:
            llm_result = llm_client.chat(messages, model=self._preferred_model())
        except RuntimeError as e:
            self._emit({"type": "error", "text": str(e)})
            return

        action_type = llm_result.get("action_type", "execute")

        # ── Handle action types ──
        if action_type == "clarify":
            question = llm_result.get("question", "請提供更多資訊")
            self._emit({"type": "clarify", "text": question})
            _update_history(sid, message, question)
            return

        if action_type == "done":
            answer = llm_result.get("answer", "")
            self._emit({"type": "result", "success": True, "text": answer, "modified": False})
            _update_history(sid, message, answer)
            return

        if action_type != "execute":
            self._emit({"type": "error", "text": "無法解析 AI 回應，請重試"})
            return

        # ── Execute ──
        tool_name   = llm_result.get("tool") or ""
        params      = llm_result.get("params") or {}
        explanation = llm_result.get("explanation") or ""

        tool_info = registry.get_tool(tool_name)
        if not tool_info:
            self._emit({
                "type": "error",
                "text": f"找不到工具「{tool_name}」，請重新描述您的需求\n可用欄位：{wb_info.get('columns', [])}",
            })
            return

        self._emit({
            "type": "tool_selected",
            "tool": tool_name,
            "name_zh": tool_info["name_zh"],
            "explanation": explanation,
        })

        if tool_info["confirm"]:
            confirm_token = str(uuid.uuid4())[:8]
            with _lock:
                SESSIONS[sid]["pending_confirm"] = {
                    "token":       confirm_token,
                    "tool":        tool_name,
                    "params":      params,
                    "wb_path":     session["workbook_path"],
                    "mode":        mode,
                    "explanation": explanation,
                }
            self._emit({
                "type":          "confirm_required",
                "confirm_token": confirm_token,
                "session_id":    sid,
                "tool":          tool_name,
                "name_zh":       tool_info["name_zh"],
                "explanation":   explanation,
                "message":       f"即將修改「{session.get('filename','工作簿')}」\n{explanation}\n確定繼續？",
            })
            return

        self._run_tool(sid, tool_name, params, wb, session["workbook_path"], mode, message, tool_info)

    def _run_tool(self, sid, tool_name, params, wb, wb_path, mode, user_msg, tool_info):
        self._emit({"type": "executing", "text": "執行中…"})
        exec_result = runner.execute_tool(
            tool_name, params, wb, wb_path, mode, backups_dir=str(BACKUPS_DIR)
        )

        if not exec_result.get("success"):
            self._emit({"type": "result", "success": False, "text": f"操作失敗：{exec_result.get('error', '未知錯誤')}"})
            _update_history(sid, user_msg, f"操作失敗：{exec_result.get('error')}")
            return

        # Register backup
        backup_path = exec_result.get("backup_path")
        if backup_path:
            tool_entry = registry.get_tool(tool_name)
            _register_backup(
                backup_path,
                SESSIONS.get(sid, {}).get("filename", "工作簿"),
                tool_name,
                tool_entry["name_zh"] if tool_entry else tool_name,
                sid,
            )

        # Save workbook if modified
        if tool_info["modifies"]:
            try:
                excel_io.save_workbook(wb, wb_path, mode)
            except Exception as e:
                self._emit({"type": "error", "text": f"儲存失敗: {e}"})
                return

        # Handle special export types (CSV)
        result_data = exec_result.get("result", {})
        export_type = result_data.get("export_type")
        dl_token = None

        if export_type == "csv":
            csv_content = result_data.get("csv_content", "")
            csv_filename = result_data.get("filename", "export.csv")
            token = str(uuid.uuid4())[:8]
            session_dir = SESSIONS_DIR / sid
            session_dir.mkdir(parents=True, exist_ok=True)
            csv_path = session_dir / csv_filename
            csv_path.write_text(csv_content, encoding="utf-8-sig")  # BOM for Excel
            DOWNLOAD_TOKENS[token] = {"path": str(csv_path), "filename": csv_filename, "content_type": "text/csv; charset=utf-8-sig"}
            dl_token = token
        elif tool_info["modifies"]:
            token = str(uuid.uuid4())[:8]
            fname = SESSIONS.get(sid, {}).get("filename", "result.xlsx")
            DOWNLOAD_TOKENS[token] = {"path": wb_path, "filename": fname, "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
            dl_token = token

        result_text = _format_result(result_data)
        self._emit({
            "type":           "result",
            "success":        True,
            "text":           result_text,
            "modified":       tool_info["modifies"],
            "download_token": dl_token,
        })
        _update_history(sid, user_msg, result_text)

    def _preferred_model(self) -> str:
        try:
            cfg_path = _ROOT / "config" / "settings.json"
            cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
            return cfg.get("agent", {}).get("default_model", "qwen2.5:7b")
        except Exception:
            return "qwen2.5:7b"

    # ── Agent: confirm pending operation ──────────────────────────────────────

    def _handle_agent_confirm(self):
        body = self.rfile.read(int(self.headers.get("Content-Length", 0)))
        try:
            data = json.loads(body)
        except Exception:
            self._json({"error": "invalid JSON"}, 400)
            return

        sid           = data.get("session_id", "")
        confirm_token = data.get("confirm_token", "")
        approved      = data.get("approved", False)

        if not self._is_safe_sid(sid):
            self._json({"error": "invalid session_id"}, 400)
            return
        session = SESSIONS.get(sid)
        if not session:
            self._json({"error": "Session 不存在"}, 404)
            return
        pending = session.get("pending_confirm")
        if not pending or pending.get("token") != confirm_token:
            self._json({"error": "無效的確認 token"}, 400)
            return

        with _lock:
            SESSIONS[sid]["pending_confirm"] = None

        if not approved:
            self._json({"status": "cancelled"})
            return

        # Set SSE headers and execute
        self.send_response(200)
        self.send_header("Content-Type", "text/event-stream; charset=utf-8")
        self.send_header("Cache-Control", "no-cache")
        self._cors()
        self.end_headers()

        try:
            wb, mode = excel_io.open_workbook(pending["wb_path"])
            tool_info = registry.get_tool(pending["tool"])
            self._run_tool(
                sid,
                pending["tool"],
                pending["params"],
                wb,
                pending["wb_path"],
                pending["mode"],
                pending.get("explanation", ""),
                tool_info,
            )
        except Exception as e:
            self._emit({"type": "error", "text": f"執行失敗: {e}"})
        finally:
            self._emit({"type": "done"})

    # ── Agent: download ───────────────────────────────────────────────────────

    def _handle_agent_download(self, token: str):
        info = DOWNLOAD_TOKENS.get(token)
        if not info:
            self._not_found()
            return
        fpath = Path(info["path"])
        if not fpath.exists():
            self._json({"error": "檔案不存在或已過期"}, 404)
            return
        data = fpath.read_bytes()
        filename = info.get("filename", "download")
        ct = info.get("content_type", "application/octet-stream")
        self.send_response(200)
        self.send_header("Content-Type", ct)
        self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        self.send_header("Content-Length", len(data))
        self._cors()
        self.end_headers()
        self.wfile.write(data)

    # ── Backup management ─────────────────────────────────────────────────────

    def _handle_backups_list(self):
        entries = _load_backup_index()
        self._json({"backups": entries, "count": len(entries)})

    def _handle_backup_restore(self):
        body = self.rfile.read(int(self.headers.get("Content-Length", 0)))
        try:
            data = json.loads(body)
        except Exception:
            self._json({"error": "invalid JSON"}, 400)
            return
        bid = data.get("backup_id", "")
        sid = data.get("session_id", "")
        entries = _load_backup_index()
        entry = next((e for e in entries if e["id"] == bid), None)
        if not entry:
            self._json({"error": f"找不到備份 ID: {bid}"}, 404)
            return
        backup_path = _ROOT / entry["backup_path"]
        if not backup_path.exists():
            self._json({"error": "備份檔不存在"}, 404)
            return
        session = SESSIONS.get(sid)
        if not session or not session.get("workbook_path"):
            self._json({"error": "Session 不存在或無工作簿"}, 404)
            return
        import shutil
        shutil.copy2(str(backup_path), session["workbook_path"])
        self._json({"status": "restored", "backup_id": bid, "timestamp": entry["timestamp"]})

    def _handle_backup_delete(self, bid: str):
        entries = _load_backup_index()
        entry = next((e for e in entries if e["id"] == bid), None)
        if not entry:
            self._json({"error": f"找不到備份 ID: {bid}"}, 404)
            return
        backup_path = _ROOT / entry["backup_path"]
        if backup_path.exists():
            backup_path.unlink()
        entries = [e for e in entries if e["id"] != bid]
        _save_backup_index(entries)
        self._json({"status": "deleted", "backup_id": bid})

    # ── Utilities ─────────────────────────────────────────────────────────────

    def _emit(self, data: dict) -> None:
        payload = f"data: {json.dumps(data, ensure_ascii=False)}\n\n"
        self.wfile.write(payload.encode("utf-8"))
        self.wfile.flush()

    def _json(self, data: dict, status: int = 200) -> None:
        payload = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", len(payload))
        self._cors()
        self.end_headers()
        self.wfile.write(payload)

    def _not_found(self) -> None:
        self.send_response(404)
        self.end_headers()
        self.wfile.write(b"Not Found")


# ─────────────────────────────────────────────────────────────────────────────

def main():
    print(f"[WinLLM] 服務啟動於 http://localhost:{PORT}")
    print(f"[WinLLM] Agent 模組: {'✓ 已啟用' if AGENT_ENABLED else '✗ 未載入'}")
    httpd = HTTPServer(("127.0.0.1", PORT), Handler)
    httpd.serve_forever()


if __name__ == "__main__":
    main()
