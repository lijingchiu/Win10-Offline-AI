"""資料清理工具 (6 個函式)"""
import io
import pandas as pd
import openpyxl
from ._base import tool


def _ws_to_df(wb, sheet: str) -> tuple:
    """Return (ws, df, headers). Raises ValueError if sheet missing."""
    name = sheet or wb.sheetnames[0]
    if name not in wb.sheetnames:
        raise ValueError(f"工作表「{name}」不存在，可用: {wb.sheetnames}")
    ws = wb[name]
    data = list(ws.values)
    if not data:
        raise ValueError(f"工作表「{name}」是空的")
    headers = [str(h) if h is not None else f"col{i}" for i, h in enumerate(data[0])]
    df = pd.DataFrame(data[1:], columns=headers)
    return ws, df, headers


def _df_to_ws(ws, df: pd.DataFrame, write_header: bool = False) -> None:
    """Overwrite ws data rows; optionally also write header row."""
    if write_header:
        for c_idx, col_name in enumerate(df.columns, start=1):
            ws.cell(row=1, column=c_idx, value=col_name)
        # Clear extra columns if df is narrower than ws
        for c_idx in range(len(df.columns) + 1, ws.max_column + 1):
            ws.cell(row=1, column=c_idx, value=None)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.value = None
    for r_idx, row in enumerate(df.itertuples(index=False, name=None), start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)


@tool(
    name_zh="刪除重複列",
    desc_zh="刪除工作表中重複的列，可指定依據哪些欄位判斷重複",
    desc_en="remove duplicate rows deduplicate keep first last",
    params_schema={
        "type": "object",
        "properties": {
            "sheet":   {"type": "string",  "description": "工作表名稱，留空用第一張"},
            "columns": {"type": "array",   "items": {"type": "string"}, "description": "依哪些欄判重複，留空則全欄比對"},
            "keep":    {"type": "string",  "enum": ["first", "last", "none"], "default": "first", "description": "保留哪筆：first/last/none"},
        },
        "required": [],
    },
    confirm=True,
    modifies=True,
)
def remove_duplicates(wb, *, sheet: str = "", columns: list = None, keep: str = "first") -> dict:
    ws, df, _ = _ws_to_df(wb, sheet)
    subset = columns or None
    if subset:
        bad = [c for c in subset if c not in df.columns]
        if bad:
            return {"error": f"欄位不存在: {bad}，可用欄位: {list(df.columns)}"}
    keep_arg = None if keep == "none" else keep
    before = len(df)
    df.drop_duplicates(subset=subset, keep=keep_arg, inplace=True)
    df.reset_index(drop=True, inplace=True)
    removed = before - len(df)
    _df_to_ws(ws, df)
    return {"message": f"已刪除 {removed} 筆重複列，剩餘 {len(df)} 筆", "removed_count": removed, "remaining_count": len(df)}


@tool(
    name_zh="去除前後空格",
    desc_zh="去除工作表欄位儲存格的前後空白字元",
    desc_en="strip whitespace trim leading trailing spaces cells",
    params_schema={
        "type": "object",
        "properties": {
            "sheet":   {"type": "string", "description": "工作表名稱，留空用第一張"},
            "columns": {"type": "array",  "items": {"type": "string"}, "description": "要處理的欄位，留空則全部文字欄"},
        },
        "required": [],
    },
    confirm=False,
    modifies=True,
)
def strip_whitespace(wb, *, sheet: str = "", columns: list = None) -> dict:
    ws, df, _ = _ws_to_df(wb, sheet)
    target = columns if columns else list(df.select_dtypes(include="object").columns)
    bad = [c for c in target if c not in df.columns]
    if bad:
        return {"error": f"欄位不存在: {bad}"}
    count = 0
    for col in target:
        mask = df[col].astype(str).str.strip() != df[col].astype(str)
        count += mask.sum()
        df[col] = df[col].apply(lambda x: x.strip() if isinstance(x, str) else x)
    _df_to_ws(ws, df)
    return {"message": f"已清理 {count} 個儲存格的前後空格", "cleaned_count": int(count)}


@tool(
    name_zh="填補空值",
    desc_zh="將工作表中的空白儲存格填入指定值或用前後值填補",
    desc_en="fill missing values null empty cells ffill bfill constant",
    params_schema={
        "type": "object",
        "properties": {
            "sheet":      {"type": "string", "description": "工作表名稱，留空用第一張"},
            "columns":    {"type": "array",  "items": {"type": "string"}, "description": "要填補的欄位，留空則全部"},
            "fill_value": {"type": "string", "description": "填補的值（method=value時使用）", "default": ""},
            "method":     {"type": "string", "enum": ["value", "ffill", "bfill"], "default": "value", "description": "填補方式：value固定值/ffill向前/bfill向後"},
        },
        "required": [],
    },
    confirm=False,
    modifies=True,
)
def fill_missing_values(wb, *, sheet: str = "", columns: list = None, fill_value: str = "", method: str = "value") -> dict:
    ws, df, _ = _ws_to_df(wb, sheet)
    target = columns if columns else list(df.columns)
    bad = [c for c in target if c not in df.columns]
    if bad:
        return {"error": f"欄位不存在: {bad}"}
    count = int(df[target].isna().sum().sum())
    if method == "ffill":
        df[target] = df[target].ffill()
    elif method == "bfill":
        df[target] = df[target].bfill()
    else:
        # Try numeric conversion; fall back to string
        for col in target:
            try:
                val = float(fill_value) if "." in str(fill_value) else int(fill_value)
            except (ValueError, TypeError):
                val = fill_value
            df[col] = df[col].fillna(val)
    _df_to_ws(ws, df)
    return {"message": f"已填補 {count} 個空白儲存格（方式：{method}）", "filled_count": count}


@tool(
    name_zh="統一欄位格式",
    desc_zh="將指定欄位的值統一轉換成日期、數字、全大寫或全小寫等格式",
    desc_en="normalize column format date number uppercase lowercase text standardize",
    params_schema={
        "type": "object",
        "properties": {
            "sheet":       {"type": "string", "description": "工作表名稱，留空用第一張"},
            "column":      {"type": "string", "description": "要轉換的欄位名稱"},
            "format_type": {"type": "string", "enum": ["date", "number", "text", "upper", "lower", "title"], "description": "轉換類型"},
            "date_format": {"type": "string", "default": "%Y-%m-%d", "description": "date 類型使用的格式字串"},
        },
        "required": ["column", "format_type"],
    },
    confirm=False,
    modifies=True,
)
def normalize_column_format(wb, *, sheet: str = "", column: str, format_type: str, date_format: str = "%Y-%m-%d") -> dict:
    ws, df, _ = _ws_to_df(wb, sheet)
    if column not in df.columns:
        return {"error": f"欄位「{column}」不存在，可用: {list(df.columns)}"}
    errors = 0
    if format_type == "date":
        def conv_date(v):
            nonlocal errors
            try:
                return pd.to_datetime(v).strftime(date_format)
            except Exception:
                errors += 1
                return v
        df[column] = df[column].apply(conv_date)
    elif format_type == "number":
        def conv_num(v):
            nonlocal errors
            try:
                return float(str(v).replace(",", "").strip())
            except Exception:
                errors += 1
                return v
        df[column] = df[column].apply(conv_num)
    elif format_type == "text":
        df[column] = df[column].astype(str).replace("nan", "")
    elif format_type == "upper":
        df[column] = df[column].apply(lambda x: str(x).upper() if pd.notna(x) else x)
    elif format_type == "lower":
        df[column] = df[column].apply(lambda x: str(x).lower() if pd.notna(x) else x)
    elif format_type == "title":
        df[column] = df[column].apply(lambda x: str(x).title() if pd.notna(x) else x)
    _df_to_ws(ws, df)
    msg = f"已將「{column}」欄轉換為 {format_type} 格式"
    if errors:
        msg += f"（{errors} 個值無法轉換，保留原值）"
    return {"message": msg, "error_count": errors}


@tool(
    name_zh="分割欄位",
    desc_zh="依分隔符號將一個欄位拆分成多個新欄位",
    desc_en="split column delimiter separate into multiple columns",
    params_schema={
        "type": "object",
        "properties": {
            "sheet":        {"type": "string", "description": "工作表名稱，留空用第一張"},
            "column":       {"type": "string", "description": "要分割的欄位"},
            "delimiter":    {"type": "string", "default": ",", "description": "分隔符號"},
            "new_columns":  {"type": "array",  "items": {"type": "string"}, "description": "新欄位名稱清單，留空自動命名"},
            "drop_original":{"type": "boolean","default": True, "description": "是否刪除原始欄位"},
        },
        "required": ["column"],
    },
    confirm=False,
    modifies=True,
)
def split_column(wb, *, sheet: str = "", column: str, delimiter: str = ",", new_columns: list = None, drop_original: bool = True) -> dict:
    ws, df, _ = _ws_to_df(wb, sheet)
    if column not in df.columns:
        return {"error": f"欄位「{column}」不存在，可用: {list(df.columns)}"}
    split_df = df[column].astype(str).str.split(delimiter, expand=True)
    n_cols = split_df.shape[1]
    if new_columns and len(new_columns) >= n_cols:
        col_names = new_columns[:n_cols]
    else:
        col_names = [f"{column}_{i+1}" for i in range(n_cols)]
    split_df.columns = col_names
    col_idx = df.columns.get_loc(column)
    for i, c in enumerate(col_names):
        df.insert(col_idx + 1 + i, c, split_df[c])
    if drop_original:
        df.drop(columns=[column], inplace=True)
    _df_to_ws(ws, df, write_header=True)
    return {"message": f"已將「{column}」分割成 {n_cols} 欄：{col_names}", "new_columns": col_names}


@tool(
    name_zh="合併欄位",
    desc_zh="將多個欄位合併成一個新欄位，可指定分隔符號",
    desc_en="merge combine columns concatenate join separator",
    params_schema={
        "type": "object",
        "properties": {
            "sheet":          {"type": "string", "description": "工作表名稱，留空用第一張"},
            "columns":        {"type": "array",  "items": {"type": "string"}, "description": "要合併的欄位清單"},
            "new_column":     {"type": "string", "description": "合併後的新欄位名稱"},
            "separator":      {"type": "string", "default": " ",  "description": "欄位間的分隔符號"},
            "drop_originals": {"type": "boolean","default": True, "description": "是否刪除原始欄位"},
        },
        "required": ["columns", "new_column"],
    },
    confirm=False,
    modifies=True,
)
def merge_columns(wb, *, sheet: str = "", columns: list, new_column: str, separator: str = " ", drop_originals: bool = True) -> dict:
    ws, df, _ = _ws_to_df(wb, sheet)
    bad = [c for c in columns if c not in df.columns]
    if bad:
        return {"error": f"欄位不存在: {bad}，可用: {list(df.columns)}"}
    df[new_column] = df[columns].astype(str).agg(separator.join, axis=1)
    if drop_originals:
        df.drop(columns=columns, inplace=True)
    _df_to_ws(ws, df, write_header=True)
    return {"message": f"已將 {columns} 合併為「{new_column}」欄（分隔符：'{separator}'）", "new_column": new_column}
