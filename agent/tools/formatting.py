"""格式化工具 (5 個函式)"""
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.cell import coordinate_from_string
from ._base import tool
from .cleaning import _ws_to_df


def _get_ws(wb, sheet: str):
    name = sheet or wb.sheetnames[0]
    if name not in wb.sheetnames:
        raise ValueError(f"工作表「{name}」不存在，可用: {wb.sheetnames}")
    return wb[name]


def _col_letter(ws, col_name: str) -> str | None:
    """Find the column letter by header value."""
    for cell in ws[1]:
        if str(cell.value) == col_name:
            return get_column_letter(cell.column)
    return None


def _hex_color(color: str) -> str:
    """Normalize color string to 6-char hex without #."""
    presets = {
        "red": "FF0000", "green": "00B050", "yellow": "FFFF00",
        "blue": "0070C0", "orange": "FFA500", "grey": "A6A6A6",
        "gray": "A6A6A6",
    }
    c = color.lstrip("#").lower()
    return presets.get(c, c.upper()[:6].zfill(6))


@tool(
    name_zh="條件格式",
    desc_zh="對指定欄位依條件套用顏色高亮（大於/小於閾值或色階）",
    desc_en="conditional format highlight color rule threshold cell",
    params_schema={
        "type": "object",
        "properties": {
            "sheet":     {"type": "string", "description": "工作表名稱"},
            "column":    {"type": "string", "description": "要套用條件格式的欄位"},
            "rule_type": {"type": "string", "enum": ["greater_than", "less_than", "equal", "color_scale"], "description": "規則類型"},
            "threshold": {"type": "number", "description": "比較閾值（color_scale 不需要）"},
            "color":     {"type": "string", "default": "red", "description": "高亮顏色：red/green/yellow/blue/orange/grey 或 16 進位"},
        },
        "required": ["column", "rule_type"],
    },
    confirm=False,
    modifies=True,
)
def apply_conditional_format(wb, *, sheet: str = "", column: str, rule_type: str, threshold=None, color: str = "red") -> dict:
    ws = _get_ws(wb, sheet)
    col_letter = _col_letter(ws, column)
    if not col_letter:
        return {"error": f"欄位「{column}」不存在"}
    max_row = ws.max_row
    cell_range = f"{col_letter}2:{col_letter}{max_row}"
    fill = PatternFill(start_color=_hex_color(color), end_color=_hex_color(color), fill_type="solid")
    if rule_type == "color_scale":
        rule = ColorScaleRule(
            start_type="min", start_color="63BE7B",
            mid_type="percentile", mid_value=50, mid_color="FFEB84",
            end_type="max", end_color="F8696B",
        )
    elif rule_type == "greater_than":
        rule = CellIsRule(operator="greaterThan", formula=[str(threshold)], fill=fill)
    elif rule_type == "less_than":
        rule = CellIsRule(operator="lessThan", formula=[str(threshold)], fill=fill)
    elif rule_type == "equal":
        rule = CellIsRule(operator="equal", formula=[str(threshold)], fill=fill)
    else:
        return {"error": f"不支援的規則類型: {rule_type}"}
    ws.conditional_formatting.add(cell_range, rule)
    return {"message": f"已對「{column}」欄套用條件格式（{rule_type}）"}


@tool(
    name_zh="調整欄寬",
    desc_zh="自動或手動調整工作表的欄位寬度",
    desc_en="column width auto fit adjust resize",
    params_schema={
        "type": "object",
        "properties": {
            "sheet":   {"type": "string", "description": "工作表名稱"},
            "columns": {"type": "array",  "items": {"type": "string"}, "description": "要調整的欄位名稱清單，留空則全部"},
            "auto":    {"type": "boolean","default": True, "description": "true=自動依內容寬度，false=使用 width 參數"},
            "width":   {"type": "number", "default": 15, "description": "固定寬度（auto=false 時使用）"},
        },
        "required": [],
    },
    confirm=False,
    modifies=True,
)
def adjust_column_width(wb, *, sheet: str = "", columns: list = None, auto: bool = True, width: float = 15) -> dict:
    ws = _get_ws(wb, sheet)
    target_letters = []
    if columns:
        for c in columns:
            letter = _col_letter(ws, c)
            if letter:
                target_letters.append(letter)
            else:
                return {"error": f"欄位「{c}」不存在"}
    else:
        from openpyxl.utils import get_column_letter
        target_letters = [get_column_letter(i) for i in range(1, ws.max_column + 1)]
    if auto:
        for letter in target_letters:
            max_len = 0
            for cell in ws[letter]:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[letter].width = min(max_len + 2, 60)
    else:
        for letter in target_letters:
            ws.column_dimensions[letter].width = width
    return {"message": f"已調整 {len(target_letters)} 個欄位的寬度（{'自動' if auto else f'{width}'}）"}


@tool(
    name_zh="設定儲存格樣式",
    desc_zh="對指定範圍的儲存格設定粗體、字型大小、背景顏色或字體顏色",
    desc_en="cell style bold font size background color format range",
    params_schema={
        "type": "object",
        "properties": {
            "sheet":      {"type": "string", "description": "工作表名稱"},
            "range_str":  {"type": "string", "description": "儲存格範圍，如 A1:D1 或 A1"},
            "bold":       {"type": "boolean","default": False},
            "font_size":  {"type": "number", "description": "字型大小"},
            "bg_color":   {"type": "string", "description": "背景顏色"},
            "font_color": {"type": "string", "description": "字體顏色"},
        },
        "required": ["range_str"],
    },
    confirm=False,
    modifies=True,
)
def set_cell_style(wb, *, sheet: str = "", range_str: str, bold: bool = False, font_size=None, bg_color: str = "", font_color: str = "") -> dict:
    ws = _get_ws(wb, sheet)
    fill = PatternFill(start_color=_hex_color(bg_color), end_color=_hex_color(bg_color), fill_type="solid") if bg_color else None
    font_kwargs: dict = {}
    if bold:
        font_kwargs["bold"] = True
    if font_size:
        font_kwargs["size"] = float(font_size)
    if font_color:
        font_kwargs["color"] = _hex_color(font_color)
    font = Font(**font_kwargs) if font_kwargs else None
    count = 0
    for row in ws[range_str]:
        cells = row if hasattr(row, "__iter__") else [row]
        for cell in cells:
            if fill:
                cell.fill = fill
            if font:
                cell.font = font
            count += 1
    return {"message": f"已套用樣式到 {count} 個儲存格（範圍：{range_str}）"}


@tool(
    name_zh="合併儲存格",
    desc_zh="合併指定範圍的儲存格",
    desc_en="merge cells range combine",
    params_schema={
        "type": "object",
        "properties": {
            "sheet":     {"type": "string", "description": "工作表名稱"},
            "range_str": {"type": "string", "description": "要合併的範圍，如 A1:C1"},
        },
        "required": ["range_str"],
    },
    confirm=False,
    modifies=True,
)
def merge_cells_range(wb, *, sheet: str = "", range_str: str) -> dict:
    ws = _get_ws(wb, sheet)
    ws.merge_cells(range_str)
    return {"message": f"已合併儲存格範圍 {range_str}"}


@tool(
    name_zh="凍結窗格",
    desc_zh="凍結工作表頂部列或左側欄，使其在捲動時保持可見",
    desc_en="freeze panes rows columns fixed header",
    params_schema={
        "type": "object",
        "properties": {
            "sheet":       {"type": "string",  "description": "工作表名稱"},
            "freeze_row":  {"type": "integer", "default": 1, "description": "凍結前幾列（0=不凍結）"},
            "freeze_col":  {"type": "integer", "default": 0, "description": "凍結前幾欄（0=不凍結）"},
        },
        "required": [],
    },
    confirm=False,
    modifies=True,
)
def freeze_panes(wb, *, sheet: str = "", freeze_row: int = 1, freeze_col: int = 0) -> dict:
    ws = _get_ws(wb, sheet)
    if freeze_row == 0 and freeze_col == 0:
        ws.freeze_panes = None
        return {"message": "已取消凍結窗格"}
    col_letter = get_column_letter(freeze_col + 1) if freeze_col > 0 else "A"
    ws.freeze_panes = f"{col_letter}{freeze_row + 1}"
    return {"message": f"已凍結前 {freeze_row} 列、前 {freeze_col} 欄"}
